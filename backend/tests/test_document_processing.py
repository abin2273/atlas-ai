from io import BytesIO
from uuid import UUID, uuid4

import pytest
from app.core.config import settings
from app.db.session import SessionLocal
from app.features.documents import processing
from app.features.documents import router as documents_router
from app.features.documents.models import Document, DocumentStatus
from app.features.documents.processing import (
    DocumentProcessingError,
    extract_document_text,
)
from app.features.documents.storage import (
    DocumentStorageError,
    get_document_storage,
)
from app.features.documents.tasks import process_document_task
from app.main import app
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfWriter


def register(client: TestClient, email: str) -> str:
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "name": "Test User", "password": "test-password-123"},
    )
    assert response.status_code == 201
    return response.json()["access_token"]


def _text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET\n".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"endstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(output)


def test_pdf_docx_and_html_parsing() -> None:
    pdf_type, pdf_text, pdf_metadata = extract_document_text(
        "manual.pdf", _text_pdf("PDF extraction works")
    )
    assert pdf_type == "application/pdf"
    assert "PDF extraction works" in pdf_text
    assert pdf_metadata["page_count"] == 1
    assert pdf_metadata["scanned_pages"] == 0

    word = WordDocument()
    word.add_paragraph("DOCX extraction works")
    word.core_properties.title = "Operations Manual"
    word.core_properties.author = "Atlas Team"
    word.core_properties.subject = "Document processing"
    word.core_properties.keywords = "atlas,search"
    word.core_properties.category = "Operations"
    word.core_properties.last_modified_by = "Atlas Editor"
    word.core_properties.revision = 3
    word_buffer = BytesIO()
    word.save(word_buffer)
    docx_type, docx_text, docx_metadata = extract_document_text(
        "manual.docx", word_buffer.getvalue()
    )
    assert (
        docx_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "DOCX extraction works" in docx_text
    assert docx_metadata["title"] == "Operations Manual"
    assert docx_metadata["author"] == "Atlas Team"
    assert docx_metadata["subject"] == "Document processing"
    assert docx_metadata["keywords"] == "atlas,search"
    assert docx_metadata["category"] == "Operations"
    assert docx_metadata["last_modified_by"] == "Atlas Editor"
    assert docx_metadata["revision"] == 3
    assert "created_at" in docx_metadata
    assert "modified_at" in docx_metadata

    html_type, html_text, html_metadata = extract_document_text(
        "manual.html",
        (
            b"<h1>Visible title</h1><script>do not index secret</script>"
            b"<p>Visible body</p>"
        ),
    )
    assert html_type == "text/html"
    assert "Visible title" in html_text and "Visible body" in html_text
    assert "do not index" not in html_text
    assert html_metadata == {}


def test_scanned_pdf_and_image_use_local_ocr(monkeypatch) -> None:
    monkeypatch.setattr(
        processing, "_run_tesseract", lambda image: "OCR extracted text"
    )
    pdf_type, pdf_text, pdf_metadata = extract_document_text("scan.pdf", _text_pdf(""))
    assert pdf_type == "application/pdf"
    assert "OCR extracted text" in pdf_text
    assert pdf_metadata["page_count"] == 1

    image_buffer = BytesIO()
    Image.new("RGB", (4, 4), color="white").save(image_buffer, format="PNG")
    image_type, image_text, image_metadata = extract_document_text(
        "scan.png", image_buffer.getvalue()
    )
    assert image_type == "image/png"
    assert image_text == "OCR extracted text"
    assert image_metadata == {
        "image_width": 4,
        "image_height": 4,
        "image_format": "PNG",
        "image_mode": "RGB",
    }

    exif_image = Image.new("RGB", (6, 5), color="white")
    exif = exif_image.getexif()
    exif[271] = "Atlas Camera"
    exif[272] = "Model X"
    exif[306] = "2026:10:01 12:00:00"
    jpeg_buffer = BytesIO()
    exif_image.save(jpeg_buffer, format="JPEG", exif=exif)
    _, _, jpeg_metadata = extract_document_text("camera.jpg", jpeg_buffer.getvalue())
    assert jpeg_metadata["image_format"] == "JPEG"
    assert jpeg_metadata["exif_make"] == "Atlas Camera"
    assert jpeg_metadata["exif_model"] == "Model X"
    assert jpeg_metadata["exif_datetime"] == "2026:10:01 12:00:00"


def test_pdf_title_and_author_metadata_are_extracted(monkeypatch) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_metadata(
        {
            "/Title": "PDF Operations",
            "/Author": "PDF Team",
            "/Subject": "Internal procedures",
            "/Keywords": "operations,atlas",
            "/Creator": "Atlas Writer",
            "/Producer": "Atlas Export",
        }
    )
    pdf_buffer = BytesIO()
    writer.write(pdf_buffer)
    monkeypatch.setattr(
        processing,
        "_ocr_image",
        lambda data: ("scanned page", {"image_width": 1, "image_height": 1}),
    )

    content_type, extracted, metadata = extract_document_text(
        "operations.pdf", pdf_buffer.getvalue()
    )
    assert content_type == "application/pdf"
    assert extracted == "scanned page"
    assert metadata["page_count"] == 1
    assert metadata["title"] == "PDF Operations"
    assert metadata["author"] == "PDF Team"
    assert metadata["subject"] == "Internal procedures"
    assert metadata["keywords"] == "operations,atlas"
    assert metadata["creator"] == "Atlas Writer"
    assert metadata["producer"] == "Atlas Export"
    assert metadata["scanned_pages"] == 1


def test_scanned_pdf_without_tesseract_returns_actionable_error(monkeypatch) -> None:
    def missing_tesseract(*args, **kwargs):
        raise processing.pytesseract.TesseractNotFoundError()

    monkeypatch.setattr(processing.pytesseract, "image_to_string", missing_tesseract)
    try:
        extract_document_text("scan.pdf", _text_pdf(""))
    except DocumentProcessingError as error:
        assert "Tesseract" in str(error) and "PATH" in str(error)
    else:
        raise AssertionError("OCR without Tesseract must fail explicitly")


def test_uploaded_document_is_queued_and_worker_processes_from_storage(
    monkeypatch,
) -> None:
    queued: list[str] = []
    monkeypatch.setattr(
        documents_router,
        "enqueue_document_processing",
        lambda document_id: queued.append(str(document_id)),
    )
    with TestClient(app) as client:
        token = register(client, f"{uuid4()}@example.com")
        headers = {"Authorization": f"Bearer {token}"}
        organization = client.post(
            "/api/v1/organizations",
            headers=headers,
            json={"name": "Queued Org", "slug": f"queued-{uuid4().hex}"},
        )
        organization_id = organization.json()["id"]
        response = client.post(
            f"/api/v1/organizations/{organization_id}/documents",
            headers=headers,
            json={
                "title": "Queued memo",
                "source_name": "queued.txt",
                "content": "queued content can be processed in a worker",
            },
        )
        assert response.status_code == 201
        document_id = response.json()["id"]
        assert queued == [document_id]
        assert response.json()["processing_status"] == "PENDING"

        with SessionLocal() as db:
            document = db.get(Document, UUID(document_id))
            assert document is not None
            assert document.raw_content is None
            assert document.storage_key
            assert (
                get_document_storage()
                .get(document.storage_key)
                .startswith(b"queued content")
            )

        assert process_document_task.run(document_id) == DocumentStatus.READY.value
        with SessionLocal() as db:
            processed = db.get(Document, UUID(document_id))
            assert processed is not None
            assert processed.processing_status == DocumentStatus.READY
            assert "queued content" in processed.content


def test_broker_failure_is_persisted_as_failed_and_reported(monkeypatch) -> None:
    def fail_enqueue(document_id):
        raise ConnectionError("Redis unavailable")

    monkeypatch.setattr(documents_router, "enqueue_document_processing", fail_enqueue)
    with TestClient(app) as client:
        token = register(client, f"{uuid4()}@example.com")
        headers = {"Authorization": f"Bearer {token}"}
        organization = client.post(
            "/api/v1/organizations",
            headers=headers,
            json={"name": "Broker failure org", "slug": f"broker-{uuid4().hex}"},
        )
        response = client.post(
            f"/api/v1/organizations/{organization.json()['id']}/documents",
            headers=headers,
            json={
                "title": "Unqueued memo",
                "source_name": "unqueued.txt",
                "content": "The broker failure must be visible",
            },
        )
        assert response.status_code == 503
        assert "could not be queued" in response.json()["detail"]
        with SessionLocal() as db:
            failed = db.query(Document).filter_by(source_name="unqueued.txt").one()
            assert failed.processing_status == DocumentStatus.FAILED
            assert failed.processing_error == "Document processing could not be queued"
            assert failed.storage_key is not None
            assert (
                get_document_storage()
                .get(failed.storage_key)
                .startswith(b"The broker failure must be visible")
            )


def test_document_delete_removes_database_record_and_stored_source() -> None:
    with TestClient(app) as client:
        token = register(client, f"{uuid4()}@example.com")
        headers = {"Authorization": f"Bearer {token}"}
        organization = client.post(
            "/api/v1/organizations",
            headers=headers,
            json={"name": "Delete Org", "slug": f"delete-{uuid4().hex}"},
        )
        organization_id = organization.json()["id"]
        created = client.post(
            f"/api/v1/organizations/{organization_id}/documents",
            headers=headers,
            json={
                "title": "Delete me",
                "source_name": "delete-me.txt",
                "content": "stored source bytes",
            },
        )
        assert created.status_code == 201
        document_id = UUID(created.json()["id"])
        with SessionLocal() as db:
            document = db.get(Document, document_id)
            assert document is not None and document.storage_key is not None
            storage_key = document.storage_key

        deleted = client.delete(
            f"/api/v1/organizations/{organization_id}/documents/{document_id}",
            headers=headers,
        )
        assert deleted.status_code == 204
        with SessionLocal() as db:
            assert db.get(Document, document_id) is None
        with pytest.raises(DocumentStorageError, match="Could not read"):
            get_document_storage().get(storage_key)


def test_upload_tracks_versions_and_persists_processing_failures() -> None:
    with TestClient(app) as client:
        token = register(client, f"{uuid4()}@example.com")
        headers = {"Authorization": f"Bearer {token}"}
        organization = client.post(
            "/api/v1/organizations",
            headers=headers,
            json={"name": "Processing Org", "slug": f"processing-{uuid4().hex}"},
        )
        organization_id = organization.json()["id"]
        endpoint = f"/api/v1/organizations/{organization_id}/documents/upload"

        first_word = WordDocument()
        first_word.add_paragraph("First version has lunar indexing")
        first_word.core_properties.title = "Lunar Operations"
        first_word.core_properties.author = "Atlas Team"
        first_buffer = BytesIO()
        first_word.save(first_buffer)
        first = client.post(
            endpoint,
            headers=headers,
            files={
                "file": (
                    "handbook.docx",
                    first_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert first.status_code == 201
        first_document = first.json()
        assert first_document["processing_status"] == "READY"
        assert first_document["version_number"] == 1
        first_metadata = first_document["extracted_metadata"]
        assert first_metadata["title"] == "Lunar Operations"
        assert first_metadata["author"] == "Atlas Team"
        assert first_metadata["revision"] == 1
        assert "created_at" in first_metadata
        assert "modified_at" in first_metadata

        second_word = WordDocument()
        second_word.add_paragraph("Second version has solar indexing")
        second_buffer = BytesIO()
        second_word.save(second_buffer)
        second = client.post(
            endpoint,
            headers=headers,
            files={
                "file": (
                    "handbook.docx",
                    second_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert second.status_code == 201
        assert second.json()["version_group_id"] == first_document["version_group_id"]
        assert second.json()["version_number"] == 2

        history = client.get(
            f"/api/v1/organizations/{organization_id}/documents/{first_document['id']}/versions",
            headers=headers,
        )
        assert history.status_code == 200
        assert [version["version_number"] for version in history.json()] == [2, 1]
        retrieval = client.get(
            f"/api/v1/organizations/{organization_id}/retrieve",
            headers=headers,
            params={"q": "solar indexing"},
        )
        assert retrieval.status_code == 200
        assert retrieval.json()["citations"][0]["version_number"] == 2

        failed = client.post(
            endpoint,
            headers=headers,
            files={"file": ("broken.pdf", b"%PDF-invalid", "application/pdf")},
        )
        assert failed.status_code == 201
        failed_document = failed.json()
        assert failed_document["processing_status"] == "FAILED"
        assert failed_document["processing_error"]

        retry = client.post(
            f"/api/v1/organizations/{organization_id}/documents/{failed_document['id']}/retry",
            headers=headers,
        )
        assert retry.status_code == 200
        assert retry.json()["processing_status"] == "FAILED"
        assert retry.json()["processing_error"]
        failed_detail = client.get(
            f"/api/v1/organizations/{organization_id}/documents/{failed_document['id']}",
            headers=headers,
        )
        assert failed_detail.status_code == 200
        assert failed_detail.json()["processing_status"] == "FAILED"
        assert failed_detail.json()["processing_error"]


def test_image_upload_runs_ocr_and_indexes_extracted_text(monkeypatch) -> None:
    def fake_ocr(image):
        return "image OCR evidence"

    monkeypatch.setattr(processing, "_run_tesseract", fake_ocr)
    image_buffer = BytesIO()
    Image.new("RGB", (8, 8), color="white").save(image_buffer, format="PNG")

    with TestClient(app) as client:
        token = register(client, f"{uuid4()}@example.com")
        headers = {"Authorization": f"Bearer {token}"}
        organization = client.post(
            "/api/v1/organizations",
            headers=headers,
            json={"name": "OCR Org", "slug": f"ocr-{uuid4().hex}"},
        )
        endpoint = f"/api/v1/organizations/{organization.json()['id']}/documents/upload"
        uploaded = client.post(
            endpoint,
            headers=headers,
            files={"file": ("scan.png", image_buffer.getvalue(), "image/png")},
        )
        assert uploaded.status_code == 201
        assert uploaded.json()["processing_status"] == "READY"
        assert uploaded.json()["extracted_metadata"] == {
            "image_width": 8,
            "image_height": 8,
            "image_format": "PNG",
            "image_mode": "RGB",
        }

        retrieved = client.get(
            f"/api/v1/organizations/{organization.json()['id']}/retrieve",
            headers=headers,
            params={"q": "OCR evidence"},
        )
        assert retrieved.status_code == 200
        assert "image OCR evidence" in retrieved.json()["citations"][0]["excerpt"]


def test_upload_rejects_unsupported_type_and_oversized_content(monkeypatch) -> None:
    monkeypatch.setattr(settings, "max_document_size_bytes", 32)
    with TestClient(app) as client:
        token = register(client, f"{uuid4()}@example.com")
        headers = {"Authorization": f"Bearer {token}"}
        organization = client.post(
            "/api/v1/organizations",
            headers=headers,
            json={"name": "Upload Limit Org", "slug": f"upload-limit-{uuid4().hex}"},
        )
        endpoint = f"/api/v1/organizations/{organization.json()['id']}/documents/upload"

        unsupported = client.post(
            endpoint,
            headers=headers,
            files={"file": ("malware.exe", b"data", "application/octet-stream")},
        )
        assert unsupported.status_code == 415

        oversized = client.post(
            endpoint,
            headers=headers,
            files={"file": ("large.txt", b"x" * 33, "text/plain")},
        )
        assert oversized.status_code == 413
