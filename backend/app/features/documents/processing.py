from contextlib import ExitStack
from html.parser import HTMLParser
from io import BytesIO
from pathlib import PurePosixPath
from typing import ClassVar
from zipfile import BadZipFile, ZipFile

import fitz
import pytesseract
from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from PIL import ExifTags, Image, UnidentifiedImageError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.config import settings

SUPPORTED_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
TYPE_EXTENSIONS = {
    content_type: extension for extension, content_type in SUPPORTED_TYPES.items()
}
TYPE_EXTENSIONS["text/x-markdown"] = ".md"
MAX_EXTRACTED_CHARACTERS = 1_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 20_000_000
MAX_OCR_PAGES = 25
MAX_OCR_IMAGE_PIXELS = 25_000_000
OCR_TIMEOUT_SECONDS = 30
MAX_METADATA_VALUE_LENGTH = 500


class DocumentProcessingError(ValueError):
    pass


class _TextHTMLParser(HTMLParser):
    _BLOCK_TAGS: ClassVar[set[str]] = {
        "address",
        "article",
        "blockquote",
        "br",
        "div",
        "dl",
        "fieldset",
        "figcaption",
        "figure",
        "footer",
        "form",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "hr",
        "li",
        "main",
        "nav",
        "ol",
        "p",
        "pre",
        "section",
        "table",
        "td",
        "th",
        "tr",
        "ul",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag in {"script", "style", "noscript"}:
            self._ignored_depth += 1
        elif not self._ignored_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif not self._ignored_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self.parts.append(data)


def normalize_document_text(text: str) -> str:
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not normalized:
        raise DocumentProcessingError("No extractable text was found in the document")
    if len(normalized) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentProcessingError(
            "Extracted text exceeds the 1 MB processing limit"
        )
    return normalized


def _extract_html(data: bytes) -> str:
    try:
        source = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentProcessingError(
            "HTML documents must use UTF-8 encoding"
        ) from error
    parser = _TextHTMLParser()
    parser.feed(source)
    parser.close()
    return normalize_document_text("".join(parser.parts))


def _metadata_text(value: object) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        text = value.isoformat()
    else:
        text = str(value).strip()
    if not text:
        return None
    return text[:MAX_METADATA_VALUE_LENGTH]


def _add_metadata(metadata: dict[str, str | int], key: str, value: object) -> None:
    if isinstance(value, bool):
        metadata[key] = int(value)
    elif isinstance(value, int):
        metadata[key] = value
    else:
        normalized = _metadata_text(value)
        if normalized is not None:
            metadata[key] = normalized


def _run_tesseract(image: Image.Image) -> str:
    try:
        return pytesseract.image_to_string(
            image,
            lang=settings.ocr_language,
            timeout=OCR_TIMEOUT_SECONDS,
        )
    except pytesseract.TesseractNotFoundError as error:
        raise DocumentProcessingError(
            "OCR requires Tesseract; install the executable and ensure it is on PATH"
        ) from error
    except (pytesseract.TesseractError, RuntimeError) as error:
        raise DocumentProcessingError("Tesseract could not OCR the document") from error


def _ocr_image(data: bytes) -> tuple[str, dict[str, str | int]]:
    try:
        with Image.open(BytesIO(data)) as image:
            if image.width * image.height > MAX_OCR_IMAGE_PIXELS:
                raise DocumentProcessingError(
                    "The uploaded image exceeds the safe OCR pixel limit"
                )
            extracted = _run_tesseract(image.convert("RGB"))
            metadata: dict[str, str | int] = {
                "image_width": image.width,
                "image_height": image.height,
            }
            _add_metadata(metadata, "image_format", image.format)
            _add_metadata(metadata, "image_mode", image.mode)
            try:
                exif = image.getexif()
                exif_tags = ExifTags.TAGS
                for tag_id, key in exif_tags.items():
                    if key in {"Make", "Model", "Orientation", "DateTime"}:
                        _add_metadata(metadata, f"exif_{key.lower()}", exif.get(tag_id))
                exif_ifd = exif.get_ifd(ExifTags.IFD.Exif)
                for tag_id, key in ExifTags.TAGS.items():
                    if key in {"DateTimeOriginal", "DateTimeDigitized"}:
                        _add_metadata(
                            metadata,
                            f"exif_{key.lower()}",
                            exif_ifd.get(tag_id),
                        )
            except (AttributeError, KeyError, TypeError, ValueError):
                pass
            return extracted, metadata
    except UnidentifiedImageError as error:
        raise DocumentProcessingError(
            "The uploaded image could not be decoded"
        ) from error
    except Image.DecompressionBombError as error:
        raise DocumentProcessingError(
            "The uploaded image exceeds the safe OCR pixel limit"
        ) from error
    except DocumentProcessingError:
        raise
    except OSError as error:
        raise DocumentProcessingError("The uploaded image could not be read") from error


def _extract_pdf(data: bytes) -> tuple[str, dict[str, str | int]]:
    if not data.startswith(b"%PDF-"):
        raise DocumentProcessingError("The uploaded file is not a valid PDF")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise DocumentProcessingError("Encrypted PDF files are not supported")
        metadata: dict[str, str | int] = {"page_count": len(reader.pages)}
        if reader.metadata:
            properties = (
                ("title", reader.metadata.title),
                ("author", reader.metadata.author),
                ("subject", reader.metadata.subject),
                ("keywords", reader.metadata.get("/Keywords")),
                ("creator", reader.metadata.get("/Creator")),
                ("producer", reader.metadata.get("/Producer")),
                ("created_at", reader.metadata.get("/CreationDate")),
                ("modified_at", reader.metadata.get("/ModDate")),
            )
            for key, value in properties:
                _add_metadata(metadata, key, value)
        page_texts: list[str] = []
        extracted_length = 0
        ocr_pages = 0
        with ExitStack() as stack:
            rendered_pdf: fitz.Document | None = None
            for page_number, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if not page_text.strip():
                    ocr_pages += 1
                    if ocr_pages > MAX_OCR_PAGES:
                        raise DocumentProcessingError(
                            f"OCR is limited to {MAX_OCR_PAGES} scanned PDF "
                            "pages per document"
                        )
                    try:
                        if rendered_pdf is None:
                            rendered_pdf = stack.enter_context(
                                fitz.open(stream=data, filetype="pdf")
                            )
                        page_to_render = rendered_pdf[page_number]
                        render_scale = 2
                        if (
                            page_to_render.rect.width
                            * page_to_render.rect.height
                            * render_scale**2
                            > MAX_OCR_IMAGE_PIXELS
                        ):
                            raise DocumentProcessingError(
                                "The scanned PDF page exceeds the safe OCR pixel limit"
                            )
                        pixmap = page_to_render.get_pixmap(
                            matrix=fitz.Matrix(render_scale, render_scale),
                            alpha=False,
                        )
                        page_text, _ = _ocr_image(pixmap.tobytes("png"))
                    except DocumentProcessingError:
                        raise
                    except (
                        fitz.FileDataError,
                        fitz.EmptyFileError,
                        RuntimeError,
                    ) as error:
                        raise DocumentProcessingError(
                            "The scanned PDF page could not be rendered for OCR"
                        ) from error
                extracted_length += len(page_text)
                if extracted_length > MAX_EXTRACTED_CHARACTERS:
                    raise DocumentProcessingError(
                        "Extracted text exceeds the 1 MB processing limit"
                    )
                page_texts.append(page_text)
    except DocumentProcessingError:
        raise
    except (PdfReadError, OSError, ValueError) as error:
        raise DocumentProcessingError("The PDF could not be parsed") from error
    metadata["scanned_pages"] = ocr_pages
    return normalize_document_text("\n".join(page_texts)), metadata


def _extract_docx(
    data: bytes,
) -> tuple[str, dict[str, str | int]]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            document_entry = archive.getinfo("word/document.xml")
            total_uncompressed_bytes = sum(
                entry.file_size for entry in archive.infolist()
            )
            has_excessive_compression = any(
                entry.compress_size > 0 and entry.file_size / entry.compress_size > 100
                for entry in archive.infolist()
            )
            if (
                document_entry.file_size > MAX_DOCX_UNCOMPRESSED_BYTES
                or total_uncompressed_bytes > MAX_DOCX_UNCOMPRESSED_BYTES
                or has_excessive_compression
            ):
                raise DocumentProcessingError(
                    "DOCX document contents exceed the processing limits"
                )
    except (BadZipFile, KeyError) as error:
        raise DocumentProcessingError(
            "The uploaded file is not a valid DOCX document"
        ) from error

    try:
        document = WordDocument(BytesIO(data))
    except (PackageNotFoundError, BadZipFile, ValueError) as error:
        raise DocumentProcessingError(
            "The DOCX document could not be parsed"
        ) from error

    metadata: dict[str, str | int] = {}
    properties = document.core_properties
    for key, value in (
        ("title", properties.title),
        ("author", properties.author),
        ("subject", properties.subject),
        ("keywords", properties.keywords),
        ("category", properties.category),
        ("comments", properties.comments),
        ("last_modified_by", properties.last_modified_by),
        ("created_at", properties.created),
        ("modified_at", properties.modified),
        ("revision", properties.revision),
    ):
        _add_metadata(metadata, key, value)
    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    return normalize_document_text("\n".join(parts)), metadata


def extract_document_text(
    filename: str,
    data: bytes,
    content_type_hint: str | None = None,
) -> tuple[str, str, dict[str, str | int]]:
    normalized_name = filename.replace("\\", "/")
    extension = PurePosixPath(normalized_name).suffix.lower()
    if extension not in SUPPORTED_TYPES and content_type_hint in TYPE_EXTENSIONS:
        extension = TYPE_EXTENSIONS[content_type_hint]
    content_type = SUPPORTED_TYPES.get(extension)
    if content_type is None:
        raise DocumentProcessingError(
            "Unsupported file type; use .txt, .md, .html, .pdf, .docx, "
            ".png, .jpg, .jpeg, .tif, or .tiff"
        )

    allowed_hints = {content_type, "application/octet-stream", "binary/octet-stream"}
    if extension == ".pdf":
        allowed_hints.add("application/x-pdf")
    if extension == ".docx":
        allowed_hints.add("application/zip")
    if extension in {".md", ".markdown"}:
        allowed_hints.add("text/plain")
        allowed_hints.add("text/x-markdown")
    if content_type_hint and content_type_hint.lower() not in allowed_hints:
        raise DocumentProcessingError(
            "File extension does not match the declared content type"
        )

    if extension in {".txt", ".md", ".markdown"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise DocumentProcessingError(
                "Text documents must use UTF-8 encoding"
            ) from error
        return content_type, normalize_document_text(text), {}
    if extension in {".html", ".htm"}:
        return content_type, _extract_html(data), {}
    if extension == ".pdf":
        extracted, metadata = _extract_pdf(data)
        return content_type, extracted, metadata
    if extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        extracted, metadata = _ocr_image(data)
        return content_type, normalize_document_text(extracted), metadata
    extracted, metadata = _extract_docx(data)
    return content_type, extracted, metadata
